#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Excel workbook for the configured self Moments export (openpyxl).

Defaults (folder / JSON / xlsx name / title) come from project/config.json when
available; WX_DELIV / WX_DATA env vars still override them.
"""
import os, json
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
try:
    import wxconfig
    _SELF = wxconfig.self_export() if wxconfig.load(required=False) is not None else {}
except SystemExit:
    _SELF = {}
DELIV = (os.environ.get("WX_DELIV") or _SELF.get("outdir")
         or os.path.join(ROOT, "deliverables"))
DATA = (os.environ.get("WX_DATA") or _SELF.get("json")
        or os.path.join(DELIV, "moments_all.json"))
d = json.load(open(DATA, encoding="utf-8"))
meta = d["meta"]
OUT = os.path.join(DELIV, meta.get("xlsx_filename")
                  or os.path.basename(_SELF.get("xlsx") or "微信朋友圈全量导出.xlsx"))
posts = sorted(d["posts"], key=lambda p: p["create_time"], reverse=True)

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALT_FILL = PatternFill("solid", fgColor="F2F5FA")
WRAP = Alignment(wrap_text=True, vertical="top")

def style_sheet(ws, widths, freeze="A2"):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
            if c.row % 2 == 0:
                c.fill = ALT_FILL

def add_sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    style_sheet(ws, widths)
    return ws

wb = Workbook()
ov = wb.active
ov.title = "概览"
ov["A1"] = meta.get("report_title") or _SELF.get("title") or "微信朋友圈全量导出（本人发布）"
ov["A1"].font = Font(bold=True, size=18)
info = [
    ("账号昵称", meta.get("account_nickname") or ""),
    ("wxid", meta.get("account_wxid") or ""),
    ("备注", meta.get("account_remark") or ""),
    ("时间范围", "%s ~ %s（UTC+8）" % (meta["range"]["start"], meta["range"]["end"])),
    ("增量导入起点", (meta["since"] + " 起（更早历史不追溯，已归档内容保留）") if meta.get("since") else "全部历史"),
    ("生成时间", meta.get("generated_at_str") or ""),
    ("朋友圈条数", meta["post_count"]),
    ("图片数量", meta["image_count"]),
    ("视频数量", meta["video_count"]),
    ("点赞总数", meta["like_count"]),
    ("评论总数", meta["comment_count"]),
    ("评论附图数量", meta["comment_media_count"]),
    ("媒体总大小", "%.1f MB" % ((meta.get("media_bytes") or 0) / 1048576.0)),
    ("历史归档保留条数（最新快照未包含）", meta.get("archived_post_count") or 0),
    ("最新快照可见条数", meta.get("visible_post_count") if meta.get("visible_post_count") is not None else "全部"),
    ("同步次数", meta.get("run_count") or ""),
    ("首次导出", meta.get("first_exported_at") or ""),
    ("最近更新", meta.get("last_updated_at") or ""),
    ("显示规则", meta.get("display_rule") or ""),
    ("数据来源", meta.get("source_db") or ""),
    ("联系人库", meta.get("contact_db") or ""),
    ("数据范围", meta.get("scope_note") or ""),
]
r = 3
for k, v in info:
    ov.cell(row=r, column=1, value=k).font = Font(bold=True)
    ov.cell(row=r, column=2, value=v)
    r += 1
ov.column_dimensions["A"].width = 16
ov.column_dimensions["B"].width = 78
for c in ov["B"]:
    c.alignment = WRAP

post_rows = []
for p in posts:
    loc = p.get("location") or {}
    post_rows.append([
        p["index"], "历史归档（最新快照未包含）" if p.get("archived") else "最新快照可见",
        p["tid"], p.get("tid_db"), p["create_time_str"], p.get("weekday") or "",
        p.get("content") or "",
        sum(1 for m in p.get("media_files") or [] if m.get("media_kind") == "image"),
        sum(1 for m in p.get("media_files") or [] if m.get("media_kind") == "video"),
        len(p["likes"]), len(p["comments"]),
        p.get("content_object_type_name") or "", loc.get("city") or "",
        loc.get("poiName") or "", loc.get("poiAddress") or "",
        loc.get("latitude") or "", loc.get("longitude") or "",
        "是" if str(p.get("private")) == "1" else "", "是" if str(p.get("isTop")) == "1" else "",
        p.get("media_dir") or ""])
ws = add_sheet(wb, "朋友圈帖子",
    ["编号", "状态", "tid(无符号)", "tid_db(有符号)", "发布时间", "星期", "正文", "图片数", "视频数", "点赞数",
     "评论数", "内容类型", "城市", "地点POI", "地址", "纬度", "经度", "私密", "置顶", "媒体目录"],
    post_rows, [6, 18, 22, 21, 19, 6, 60, 7, 7, 7, 7, 10, 10, 26, 24, 11, 11, 6, 6, 30])
for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
    row[0].alignment = WRAP

like_rows = []
for p in posts:
    for i, l in enumerate(p["likes"], 1):
        like_rows.append([p["index"], p["create_time_str"], (p.get("content") or "")[:30], i,
                          l.get("remark") or "", l.get("wechat_nickname") or "",
                          l.get("display_name") or l.get("wechat_nickname") or "",
                          l.get("notice_nickname") or "", l.get("create_time_str") or "",
                          l.get("username") or ""])
add_sheet(wb, "点赞",
    ["帖子编号", "帖子时间", "正文摘要", "序号", "备注", "微信名", "显示名（备注（微信名））",
     "通知昵称", "点赞时间", "username"],
    like_rows, [8, 19, 24, 5, 16, 18, 26, 16, 19, 24])

com_rows = []
for p in posts:
    for i, c in enumerate(p["comments"], 1):
        rr = c.get("ref_comment") or {}
        com_rows.append([p["index"], p["create_time_str"], (p.get("content") or "")[:30], i,
                         c.get("comment_id") or "", c.get("remark") or "",
                         c.get("wechat_nickname") or "", c.get("display_name") or "",
                         c.get("notice_nickname") or "", c.get("content") or "",
                         c.get("create_time_str") or "", c.get("ref_comment_id") or "",
                         rr.get("display_name") or "", (rr.get("content") or "")[:120],
                         c.get("comment_imageinfo_count") or "0", c.get("username") or ""])
ws = add_sheet(wb, "评论",
    ["帖子编号", "帖子时间", "正文摘要", "序号", "comment_id", "备注", "微信名",
     "显示名（备注（微信名））", "通知昵称", "评论内容", "评论时间", "回复评论ID", "回复对象显示名",
     "被回复内容摘要", "图片数", "username"],
    com_rows, [8, 19, 24, 5, 8, 16, 18, 26, 16, 50, 19, 10, 24, 40, 7, 24])
for row in ws.iter_rows(min_row=2, min_col=10, max_col=10):
    row[0].alignment = WRAP

img_rows, vid_rows = [], []
for p in posts:
    for m in p.get("media_files") or []:
        row = [p["index"], p["create_time_str"], (p.get("content") or "")[:30], m.get("media_index"),
               os.path.basename(m.get("file") or ""), m.get("decrypted_bytes") or "",
               m.get("format") or "", m.get("width_xml") or "", m.get("height_xml") or "",
               m.get("total_size_xml") or "", m.get("status") or "", m.get("file") or ""]
        if m.get("media_kind") == "video":
            vid_rows.append(row[:7] + [m.get("video_duration") or "", m.get("thumb_file") or ""] + row[7:9] + [m.get("status") or "", m.get("file") or ""])
        else:
            img_rows.append(row)
add_sheet(wb, "图片",
    ["帖子编号", "帖子时间", "正文摘要", "媒体序号", "文件名", "解密后字节", "格式", "宽", "高",
     "XML总大小", "状态", "相对路径"],
    img_rows, [8, 19, 24, 7, 16, 12, 6, 7, 7, 11, 8, 34])
add_sheet(wb, "视频",
    ["帖子编号", "帖子时间", "正文摘要", "媒体序号", "文件名", "解析后字节", "格式", "视频时长",
     "封面文件", "宽", "高", "状态", "相对路径"],
    vid_rows, [8, 19, 24, 7, 16, 12, 6, 9, 20, 7, 7, 8, 34])

cm_rows = []
for p in posts:
    for m in (p.get("comment_media_files") or []):
        cm_rows.append([p["index"], p["create_time_str"], m.get("comment_id"), m.get("comment_media_index"),
                        os.path.basename(m.get("file") or ""), m.get("status") or "",
                        m.get("width_xml") or "", m.get("height_xml") or "", m.get("file") or ""])
add_sheet(wb, "评论附图",
    ["帖子编号", "帖子时间", "评论ID", "评论图序号", "文件名", "状态", "宽", "高", "相对路径"],
    cm_rows, [8, 19, 12, 10, 16, 10, 7, 7, 34])

cnt = Counter()
for p in posts:
    for l in p["likes"]:
        cnt[l.get("display_name") or l.get("wechat_nickname") or ""] += 1
rank = [[i, n, k] for i, (k, n) in enumerate(cnt.most_common(), 1)]
add_sheet(wb, "点赞排行", ["排名", "点赞次数", "好友（备注（微信名））"], rank, [6, 10, 40])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("xlsx:", OUT, os.path.getsize(OUT) // 1024, "KB")
