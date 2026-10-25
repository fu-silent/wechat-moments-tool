#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify the self-deliverable folder is complete and internally consistent.

Expected values come from `project/config.json` -> `self_export.expected`
(optional; leave null to skip historical-count assertions). File names and
folder come from `self_export.*` as well.

Checks:
  * JSON meta counts vs lists/posts (always)
  * optional historical expectations from config (if provided)
  * every media / raw XML / thumb reference resolves
  * sha256 spot-check against recorded values
  * no stale "data/all/" paths anywhere
  * HTML image refs, card + sidebar link counts
  * CSV row counts
  * XLSX sheets and row counts
  * zip integrity for zips present
Exit code 0 = all checks pass.
"""
import csv
import hashlib
import json
import os
import random
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import wxconfig

ROOT = wxconfig.ROOT
SELF = wxconfig.self_export()
DELIV = SELF["outdir"]
JSON_PATH = SELF["json"]
errors, notes = [], []


def err(msg):
    errors.append(msg)


def sha256_file(p, limit=1 << 20):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(limit), b""):
            h.update(chunk)
    return h.hexdigest()


if not os.path.isdir(DELIV):
    print("FAIL: deliverable dir not found:", DELIV)
    sys.exit(1)
if not os.path.isfile(JSON_PATH):
    print("FAIL: JSON not found:", JSON_PATH)
    sys.exit(1)

J = json.load(open(JSON_PATH, encoding="utf-8"))
posts, media, cmedia, meta = J["posts"], J["media"], J["comment_media"], J["meta"]

# ---- 1. counts: meta vs recomputed (always) + optional config expectations ----
recomputed = {
    "post_count": len(posts),
    "media_count": len(media),
    "image_count": sum(1 for m in media if m.get("media_kind") == "image"),
    "video_count": sum(1 for m in media if m.get("media_kind") == "video"),
    "comment_media_count": len(cmedia),
    "like_count": sum(len(p.get("likes") or []) for p in posts),
    "comment_count": sum(len(p.get("comments") or []) for p in posts),
}
for key, got in recomputed.items():
    want = meta.get(key)
    if want != got:
        err("meta.%s=%r but recomputed=%r" % (key, want, got))
expected = SELF.get("expected") or {}
for key, want in expected.items():
    if key not in recomputed:
        notes.append("config self_export.expected.%s ignored (not a known metric)" % key)
        continue
    if recomputed[key] != want:
        err("%s=%r, expected %r (from config self_export.expected)" % (key, recomputed[key], want))
ok_media = [m for m in media if m.get("status") == "ok"]
ext_media = [m for m in media if m.get("status") == "external"]
if meta.get("media_downloaded") != len(ok_media):
    err("media_downloaded mismatch: %r vs %r" % (meta.get("media_downloaded"), len(ok_media)))
if meta.get("media_failed") != len(media) - len(ok_media) - len(ext_media):
    err("media_failed mismatch: %r" % meta.get("media_failed"))
notes.append("counts: posts=%d media=%d(ok=%d,missing=%d,external=%d) comment_media=%d likes=%d comments=%d"
             % (len(posts), len(media), len(ok_media),
                len(media) - len(ok_media) - len(ext_media), len(ext_media),
                len(cmedia), meta["like_count"], meta["comment_count"]))
if expected:
    notes.append("config expectations checked: " + ", ".join(sorted(expected)))

# ---- 2. references exist ----
missing = []
refs = 0


def check_path(p):
    global refs
    if not p:
        return
    refs += 1
    full = os.path.join(DELIV, p)
    if not (os.path.isfile(full) and os.path.getsize(full) > 0):
        missing.append(p)


for p in posts:
    check_path(p.get("raw_xml_file"))
    for r in p.get("media_files") or []:
        check_path(r.get("file"))
        check_path(r.get("thumb_file"))
        check_path((r.get("thumb") or {}).get("file"))
    for r in p.get("comment_media_files") or []:
        check_path(r.get("file"))
for r in cmedia:
    check_path(r.get("file"))
if missing:
    err("%d referenced files missing (e.g. %r)" % (len(missing), missing[:5]))
notes.append("refs checked=%d missing=%d" % (refs, len(missing)))

# media dir should have no unreferenced image/video files
referenced = set()
for p in posts:
    for r in p.get("media_files") or []:
        if r.get("file"):
            referenced.add(os.path.normpath(r["file"]))
        if r.get("thumb_file"):
            referenced.add(os.path.normpath(r["thumb_file"]))
        if (r.get("thumb") or {}).get("file"):
            referenced.add(os.path.normpath(r["thumb"]["file"]))
    for r in p.get("comment_media_files") or []:
        if r.get("file"):
            referenced.add(os.path.normpath(r["file"]))
for r in cmedia:
    if r.get("file"):
        referenced.add(os.path.normpath(r["file"]))
on_disk = set()
for dirpath, _dirs, files in os.walk(os.path.join(DELIV, "media")):
    for f in files:
        full = os.path.join(dirpath, f)
        on_disk.add(os.path.normpath(os.path.relpath(full, DELIV)))
extra = on_disk - referenced
if extra:
    err("%d unreferenced files under media/ (e.g. %r)" % (len(extra), sorted(extra)[:5]))
notes.append("media files on disk=%d referenced=%d extra=%d" % (len(on_disk), len(referenced), len(extra)))
n_thumb = sum(1 for m in media if m.get("media_kind") == "video" and m.get("thumb_file"))
if meta.get("video_thumb_files") is not None and n_thumb != meta.get("video_thumb_files"):
    err("video thumb files=%r but meta.video_thumb_files=%r" % (n_thumb, meta.get("video_thumb_files")))
notes.append("video cover thumbs recorded: %d" % n_thumb)

# ---- 3. sha256 spot check ----
random.seed(20261019)
sample = random.sample([m for m in ok_media if m.get("file") and m.get("sha256")], min(12, len(ok_media)))
bad_sha = []
for m in sample:
    if sha256_file(os.path.join(DELIV, m["file"])) != m["sha256"]:
        bad_sha.append(m["file"])
if bad_sha:
    err("sha256 mismatch: %r" % bad_sha)
notes.append("sha256 spot-check: %d/%d ok" % (len(sample) - len(bad_sha), len(sample)))

# ---- 4. stale paths ----
names = [os.path.basename(JSON_PATH), os.path.basename(SELF["md"]), os.path.basename(SELF["html"])]
for name in names:
    p = os.path.join(DELIV, name)
    if os.path.exists(p):
        t = open(p, encoding="utf-8", errors="replace").read()
        for pat in ("data/all/", "data/window/", "report/moments"):
            if pat in t:
                err("%s still contains %r" % (name, pat))
for name in ("posts.csv", "media.csv"):
    p = os.path.join(DELIV, "csv", name)
    if os.path.exists(p):
        t = open(p, encoding="utf-8", errors="replace").read()
        if "data/all/" in t:
            err("%s still contains data/all/" % name)

# ---- 5. HTML ----
html_path = SELF["html"]
if not os.path.isfile(html_path):
    err("html missing: %s" % SELF["html"])
else:
    H = open(html_path, encoding="utf-8", errors="replace").read()
    img_srcs = re.findall(r'<img[^>]+src="([^"]+)"', H)
    bad_img = [s for s in img_srcs if not os.path.isfile(os.path.join(DELIV, s))]
    if bad_img:
        err("%d html img src missing (e.g. %r)" % (len(bad_img), bad_img[:5]))
    post_anchors = len(re.findall(r'id="post-\d+"', H))
    side_links = len(re.findall(r'data-text="', H))
    jump_links = len(re.findall(r'href="#post-\d+"', H))
    if post_anchors != len(posts):
        err("html post cards=%d, expected %d" % (post_anchors, len(posts)))
    if side_links != len(posts):
        err("html sidebar links=%d, expected %d" % (side_links, len(posts)))
    title = meta.get("report_title") or SELF["title"]
    if title and ("<title>" + title) not in H:
        err("html <title> does not contain report title %r" % title)
    if "数据完整性说明" not in H:
        err("html missing data-integrity notice")
    if "repeat(3,minmax(0,1fr))" not in H:
        err("html photo grid is not 3 columns")
    notes.append("html: %d bytes, imgs=%d (missing=%d), cards=%d, sidebar links=%d, total post jumps=%d"
                 % (len(H), len(img_srcs), len(bad_img), post_anchors, side_links, jump_links))

# ---- 6. CSV ----
csv_expected = {
    "posts.csv": len(posts),
    "likes.csv": meta["like_count"],
    "comments.csv": meta["comment_count"],
    "media.csv": len(media),
    "comment_media.csv": len(cmedia),
}
for name, want in csv_expected.items():
    p = os.path.join(DELIV, "csv", name)
    if not os.path.exists(p):
        err("csv missing: " + name)
        continue
    with open(p, encoding="utf-8-sig", newline="") as f:
        rows = sum(1 for _ in csv.reader(f)) - 1
    if rows != want:
        err("csv %s rows=%d expected %d" % (name, rows, want))
notes.append("csv rows: " + ", ".join("%s=%d" % (k, v) for k, v in csv_expected.items()))

# ---- 7. XLSX ----
xlsx = os.path.join(DELIV, os.path.basename(SELF["xlsx"]))
if not os.path.isfile(xlsx):
    err("xlsx missing: %s" % SELF["xlsx"])
else:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, read_only=True)
        sheets = {ws.title: ws.max_row - 1 for ws in wb.worksheets}
        wb.close()
        want_sheets = {"朋友圈帖子": len(posts), "点赞": meta["like_count"],
                       "评论": meta["comment_count"], "图片": meta["image_count"],
                       "视频": meta["video_count"], "评论附图": len(cmedia)}
        for title, want in want_sheets.items():
            got = sheets.get(title)
            if got != want:
                err("xlsx sheet %r rows=%r expected %r" % (title, got, want))
        notes.append("xlsx sheets: " + ", ".join("%s=%d" % (k, v) for k, v in sheets.items()))
    except Exception as e:  # noqa: BLE001
        err("xlsx check failed: %r" % (e,))

# ---- 8. zip (optional) ----
zs = [os.path.join(DELIV, f) for f in os.listdir(DELIV) if f.endswith(".zip")]
zs += [os.path.join(ROOT, f) for f in os.listdir(ROOT) if f.endswith(".zip")]
for zp in zs:
    try:
        with zipfile.ZipFile(zp) as z:
            bad = z.testzip()
            n = len(z.namelist())
        if bad:
            err("zip corrupt at %r" % bad)
        notes.append("zip: %s %.1fMB entries=%d" % (os.path.basename(zp), os.path.getsize(zp) / 1e6, n))
    except Exception as e:  # noqa: BLE001
        err("zip check failed: %r" % (e,))

# ---- report ----
print("\n".join("  - " + n for n in notes))
if errors:
    print("\nFAIL (%d):" % len(errors))
    for e in errors:
        print("  [x] " + e)
    sys.exit(1)
print("\nPASS: deliverables verified (%d checks, 0 errors)" % (len(notes) + 1))
