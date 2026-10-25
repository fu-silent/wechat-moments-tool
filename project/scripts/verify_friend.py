#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify an independent friend-Moments deliverable directory (export_friend.py output).

Usage:
  python3 verify_friend.py --dir deliverables_<name> [--json moments_<name>.json]
                           [--zip wechat-moments-<name>-export-YYYY-MM-DD.zip]

Exit code 0 = all checks pass.
"""
import argparse, csv, hashlib, json, os, random, re, sys, zipfile

ap = argparse.ArgumentParser()
ap.add_argument("--dir", required=True)
ap.add_argument("--json", default="")
ap.add_argument("--zip", default="")
a = ap.parse_args()
DELIV = os.path.abspath(a.dir)
errors, notes = [], []

def err(m):
    errors.append(m)

def sha256_file(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

if not os.path.isdir(DELIV):
    print("FAIL: dir not found", DELIV); sys.exit(1)
json_path = os.path.join(DELIV, a.json) if a.json else ""
if not json_path or not os.path.exists(json_path):
    cands = sorted(x for x in os.listdir(DELIV) if x.endswith(".json"))
    if not cands:
        print("FAIL: no JSON in", DELIV); sys.exit(1)
    json_path = os.path.join(DELIV, cands[0])
J = json.load(open(json_path, encoding="utf-8"))
meta = J["meta"]; posts, media, cmedia = J["posts"], J["media"], J["comment_media"]
slug = os.path.splitext(os.path.basename(json_path))[0]
html_name = meta.get("html_filename") or (slug + ".html")
md_name = meta.get("md_filename") or (slug + ".md")
xlsx_name = meta.get("xlsx_filename") or ""
TITLE = meta.get("report_title") or ""

print("verify:", DELIV)
print("json:", os.path.basename(json_path), "| title:", TITLE)

# ---- 1. counts / identity ----
checks = [
    ("post_count", len(posts), meta.get("post_count")),
    ("media_count", len(media), meta.get("media_count")),
    ("image_count", sum(1 for m in media if m.get("media_kind") == "image"), meta.get("image_count")),
    ("video_count", sum(1 for m in media if m.get("media_kind") == "video"), meta.get("video_count")),
    ("comment_media_count", len(cmedia), meta.get("comment_media_count")),
    ("like_count", sum(len(p.get("likes") or []) for p in posts), meta.get("like_count")),
    ("comment_count", sum(len(p.get("comments") or []) for p in posts), meta.get("comment_count")),
]
for name, got, want in checks:
    if got != want:
        err("%s recomputed=%r meta=%r" % (name, got, want))
if meta.get("image_count", 0) + meta.get("video_count", 0) != meta.get("media_count", 0):
    err("image+video != media_count")
if len({p["index"] for p in posts}) != len(posts) or sorted(p["index"] for p in posts) != list(range(1, len(posts) + 1)):
    err("post indexes not 1..N unique")
wx = meta.get("target_wxid")
if not wx:
    err("meta.target_wxid missing")
bad_user = [p["index"] for p in posts if p.get("username") != wx]
if bad_user:
    err("posts not from target wxid: %r" % bad_user[:10])
ok_media = [m for m in media if m.get("status") == "ok"]
ext_media = [m for m in media if m.get("status") == "external"]
if meta.get("media_downloaded") != len(ok_media):
    err("media_downloaded %r vs %r" % (meta.get("media_downloaded"), len(ok_media)))
if meta.get("media_failed") != len(media) - len(ok_media) - len(ext_media):
    err("media_failed %r vs %r" % (meta.get("media_failed"), len(media) - len(ok_media) - len(ext_media)))
ok_cm = [m for m in cmedia if m.get("status") == "ok"]
if meta.get("comment_media_downloaded") != len(ok_cm):
    err("comment_media_downloaded %r vs %r" % (meta.get("comment_media_downloaded"), len(ok_cm)))
notes.append("counts: posts=%d media=%d(ok=%d,missing=%d,external=%d) comment_media=%d(ok=%d) likes=%d comments=%d"
             % (len(posts), len(media), len(ok_media), len(media) - len(ok_media) - len(ext_media), len(ext_media),
                len(cmedia), len(ok_cm), meta["like_count"], meta["comment_count"]))

# ---- 2. references ----
missing, refs = [], 0
def check_path(rel):
    global refs
    if not rel:
        return
    refs += 1
    if os.path.isabs(rel):
        err("absolute path in JSON: %r" % rel)
        return
    full = os.path.join(DELIV, rel)
    if not (os.path.isfile(full) and os.path.getsize(full) > 0):
        missing.append(rel)

referenced = set()
for p in posts:
    check_path(p.get("raw_xml_file"))
    if p.get("raw_xml_file"):
        referenced.add(os.path.normpath(p["raw_xml_file"]))
    for r in p.get("media_files") or []:
        check_path(r.get("file")); check_path(r.get("thumb_file")); check_path((r.get("thumb") or {}).get("file"))
        for key in ("file", "thumb_file"):
            if r.get(key):
                referenced.add(os.path.normpath(r[key]))
        if (r.get("thumb") or {}).get("file"):
            referenced.add(os.path.normpath(r["thumb"]["file"]))
    for r in p.get("comment_media_files") or []:
        check_path(r.get("file"))
        if r.get("file"):
            referenced.add(os.path.normpath(r["file"]))
for r in cmedia:
    check_path(r.get("file"))
    if r.get("file"):
        referenced.add(os.path.normpath(r["file"]))
if missing:
    err("%d referenced files missing (e.g. %r)" % (len(missing), missing[:5]))
on_disk = set()
for dirpath, _dirs, files in os.walk(os.path.join(DELIV, "media")):
    for f in files:
        on_disk.add(os.path.normpath(os.path.relpath(os.path.join(dirpath, f), DELIV)))
for dirpath, _dirs, files in os.walk(os.path.join(DELIV, "raw_xmls")):
    for f in files:
        on_disk.add(os.path.normpath(os.path.relpath(os.path.join(dirpath, f), DELIV)))
extra = on_disk - referenced
if extra:
    err("%d unreferenced files under media/ raw_xmls/ (e.g. %r)" % (len(extra), sorted(extra)[:5]))
notes.append("refs checked=%d missing=%d | files on disk=%d extra=%d" % (refs, len(missing), len(on_disk), len(extra)))

# ---- 3. sha256 spot-check ----
random.seed(20261019)
sample = random.sample([m for m in ok_media if m.get("file") and m.get("sha256")], min(12, len(ok_media))) if ok_media else []
bad_sha = [m["file"] for m in sample if sha256_file(os.path.join(DELIV, m["file"])) != m["sha256"]]
if bad_sha:
    err("sha256 mismatch: %r" % bad_sha)
notes.append("sha256 spot-check: %d/%d ok" % (len(sample) - len(bad_sha), len(sample)))

# ---- 4. HTML / MD ----
hp = os.path.join(DELIV, html_name)
if not os.path.exists(hp):
    err("html missing: %s" % html_name)
else:
    H = open(hp, encoding="utf-8", errors="replace").read()
    if "<title>" + TITLE not in H:
        err("html <title> does not contain report title")
    if TITLE and ("<h1>" + TITLE) not in H:
        err("html <h1> does not equal report title")
    img_srcs = re.findall(r'<img[^>]+src="([^"]+)"', H)
    bad_img = [s for s in img_srcs if not (os.path.isfile(os.path.join(DELIV, s)) or os.path.isfile(os.path.join(os.path.dirname(hp), s)))]
    if bad_img:
        err("%d html img src missing (e.g. %r)" % (len(bad_img), bad_img[:5]))
    if 'id="post-' not in H:
        err("html has no post cards")
    cards = len(re.findall(r'id="post-\d+"', H))
    side = len(re.findall(r'data-text="', H))
    if cards != len(posts):
        err("html cards=%d expected %d" % (cards, len(posts)))
    if side != len(posts):
        err("html sidebar links=%d expected %d" % (side, len(posts)))
    if "repeat(3,minmax(0,1fr))" not in H:
        err("html photo grid not 3 columns")
    if meta.get("header_account") and meta["header_account"].split(" · ")[0] not in H:
        err("html header_account line missing")
    if len(media) - len(ok_media) - len(ext_media) > 0 and "数据完整性说明" not in H:
        err("html has failed media but no integrity notice")
    for bad in ("moments_all", "deliverables/", "data/all/", "微信朋友圈全量导出"):
        if bad in H:
            err("html contains leftover %r" % bad)
    notes.append("html: %d bytes, imgs=%d(missing=%d), cards=%d, side links=%d"
                 % (len(H), len(img_srcs), len(bad_img), cards, side))
mp = os.path.join(DELIV, md_name)
if os.path.exists(mp):
    M = open(mp, encoding="utf-8", errors="replace").read()
    if TITLE not in M.split("\n")[0]:
        err("markdown first line title mismatch")
    if "moments_all" in M:
        err("markdown contains leftover moments_all")
    notes.append("md: %d bytes" % len(M))
else:
    err("markdown missing: %s" % md_name)

# ---- 5. CSV ----
want = {"posts.csv": len(posts), "likes.csv": meta["like_count"], "comments.csv": meta["comment_count"],
        "media.csv": len(media), "comment_media.csv": len(cmedia)}
for name, n in want.items():
    p = os.path.join(DELIV, "csv", name)
    if not os.path.exists(p):
        err("csv missing: " + name); continue
    with open(p, encoding="utf-8-sig", newline="") as f:
        rows = sum(1 for _ in csv.reader(f)) - 1
    if rows != n:
        err("csv %s rows=%d expected %d" % (name, rows, n))
notes.append("csv: " + ", ".join("%s=%d" % (k, v) for k, v in want.items()))

# ---- 6. XLSX ----
if xlsx_name:
    xlsx = os.path.join(DELIV, xlsx_name)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, read_only=True)
        sheets = {ws.title: ws.max_row - 1 for ws in wb.worksheets}
        wb.close()
        exp = {"朋友圈帖子": len(posts), "点赞": meta["like_count"], "评论": meta["comment_count"],
               "图片": meta["image_count"], "视频": meta["video_count"], "评论附图": len(cmedia)}
        for t, n in exp.items():
            if sheets.get(t) != n:
                err("xlsx sheet %r rows=%r expected %r" % (t, sheets.get(t), n))
        from openpyxl import load_workbook as lw
        wb2 = lw(xlsx)
        a1 = wb2["概览"]["A1"].value
        wb2.close()
        if a1 != TITLE:
            err("xlsx A1=%r expected %r" % (a1, TITLE))
        notes.append("xlsx: " + ", ".join("%s=%d" % (k, v) for k, v in sorted(sheets.items())))
    except Exception as e:
        err("xlsx check failed: %r" % (e,))

# ---- 7. zip (optional) ----
zs = [a.zip] if a.zip else [os.path.join(DELIV, f) for f in os.listdir(DELIV) if f.endswith(".zip")]
if zs:
    zp = zs[0]
    try:
        with zipfile.ZipFile(zp) as z:
            bad = z.testzip(); n = len(z.namelist())
        if bad:
            err("zip corrupt at %r" % bad)
        notes.append("zip: %s %.1fMB entries=%d" % (os.path.basename(zp), os.path.getsize(zp) / 1e6, n))
    except Exception as e:
        err("zip check failed: %r" % (e,))

print("\n".join("  - " + n for n in notes))
if errors:
    print("\nFAIL (%d):" % len(errors))
    for e in errors:
        print("  [x] " + e)
    sys.exit(1)
print("\nPASS: friend deliverables verified (%d checks, 0 errors)" % (len(notes) + 1))
